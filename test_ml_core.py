"""Pengujian cepat fungsi utama tanpa menjalankan server Streamlit.

Jalankan: python test_ml_core.py
"""
from io import BytesIO
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from ml_core import (
    DEFAULT_TARGET, create_excel_report,
    predict_customer_profile, run_logistic_analysis,
)

DIR = Path(__file__).resolve().parent
df = pd.read_csv(DIR / 'dataset_pelanggan.csv')
features = [c for c in df.columns if c not in (DEFAULT_TARGET, 'ID_Pelanggan')]

def fit(source: pd.DataFrame, threshold: float = .5):
    return run_logistic_analysis(source, DEFAULT_TARGET, 'Berhenti', features,
                                 train_percent=80, threshold=threshold, C=1.0,
                                 balanced=False, random_seed=42)

r = fit(df)
a = r['metrics']
assert df.shape == (1000, 10)
assert df[DEFAULT_TARGET].value_counts().to_dict() == {'Bertahan': 699, 'Berhenti': 301}
assert (a['TP'], a['FN'], a['TN'], a['FP']) == (31, 29, 128, 12)
assert (a['N Training'], a['N Testing']) == (800, 200)
assert abs(a['Recall'] - 31/60) < 1e-12
assert set(r['training']['ID_Pelanggan']).isdisjoint(set(r['testing']['ID_Pelanggan']))
assert len(r['full']) == 1000
assert (r['testing']['Kategori_Confusion_Matrix'] == 'FN').sum() == a['FN']
assert (r['testing']['Kategori_Confusion_Matrix'] == 'TP').sum() == a['TP']
# Excel sumber yang diunggah dapat dipakai pada app via sheet DATA_PELANGGAN.
excel = pd.read_excel(DIR / 'Dataset_Bab_2_Customer_Churn_1000_Pelanggan.xlsx',
                      sheet_name='DATA_PELANGGAN')
assert excel.shape == df.shape
assert excel.astype(str).equals(df.astype(str))
r2 = fit(excel)
assert r2['metrics'] == a
assert np.allclose(r['scores'], r2['scores'])
# Threshold berubah: hasil probabilitas dasar sama, kelas dapat berbeda.
r_low = fit(df, .3)
assert np.allclose(r_low['scores'], r['scores'])
assert r_low['metrics']['Recall'] >= a['Recall']
# Prediksi satu profil dan laporan Excel.
profile = {}
for col in features:
    if col in r['numeric']:
        profile[col] = float(r['X_train'][col].median())
    else:
        profile[col] = 'Basic'
profile['Kepuasan_Pelanggan'] = 3
p = predict_customer_profile(r, profile)
assert 0 <= p['probability'] <= 1
assert p['prediction'] in ('Bertahan', 'Berhenti')
wb = load_workbook(BytesIO(create_excel_report(r, p)), read_only=True, data_only=True)
assert len(wb.sheetnames) == 11
assert wb['Ringkasan']['B5'].value == a['TP']
assert wb['Ringkasan']['B6'].value == a['FN']
assert abs(wb['Ringkasan']['B7'].value - a['Recall']) < 1e-10
assert abs(wb['Ringkasan']['B8'].value - p['probability']) < 1e-10
assert wb['Data Training'].max_row == 801
assert wb['Data Testing'].max_row == 201
assert wb['Prediksi Testing'].max_row == 201
assert wb['Simulasi Profil'].max_row >= len(profile) + 2
# Sebelum Q4 dijalankan: file tetap bisa diunduh, Q4 ditandai belum ada.
wb2 = load_workbook(BytesIO(create_excel_report(r)), read_only=True, data_only=True)
assert wb2['Ringkasan']['B8'].value == 'Belum dijalankan'
print('SEMUA TES LULUS: dataset Excel/CSV, 800/200, 4 pertanyaan, determinisme, ekspor 11 worksheet.')
print('TP:',a['TP'],'FN:',a['FN'],'Recall:',f"{a['Recall']:.1%}",'Contoh probabilitas:',f"{p['probability']:.1%}")
